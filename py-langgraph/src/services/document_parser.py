"""有界文档解析器，支持文本、PDF 和 DOCX。"""

from __future__ import annotations

import io
import uuid
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from pypdf import PdfReader


SESSION_DOCUMENT_EXTENSIONS = {".txt", ".md", ".markdown", ".json", ".csv", ".pdf", ".docx"}
SESSION_DOCUMENT_MAX_BYTES = 2_500_000
SESSION_DOCUMENT_MAX_TEXT_BYTES = 1_000_000
SESSION_DOCUMENT_MAX_PART_CHARS = 20_000
SESSION_DOCUMENT_JSON_MAX_BYTES = 5_000_000


class DocumentParseError(ValueError):
    """表示用户上传文档无法通过格式、大小或正文校验。"""

    # 初始化文档解析错误及其对外 HTTP 状态。
    def __init__(self, code: str, message: str, status_code: int):
        self.code = code
        self.status_code = status_code
        super().__init__(message)


# 将文本正文切分为有界的临时上下文段落。
def _build_parts(content: str, page: int | None = None) -> list[dict[str, Any]]:
    normalized = content.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return []
    return [
        {
            "partIndex": index,
            "page": page,
            "content": chunk.strip()[:SESSION_DOCUMENT_MAX_PART_CHARS],
        }
        for index, chunk in enumerate(normalized.split("\n\n"))
        if chunk.strip()
    ]


# 从 PDF 每一页提取文本并保留页码。
def _parse_pdf(buffer: bytes) -> list[dict[str, Any]]:
    try:
        reader = PdfReader(io.BytesIO(buffer))
        parts: list[dict[str, Any]] = []
        for page_number, page in enumerate(reader.pages, 1):
            page_text = page.extract_text() or ""
            for part in _build_parts(page_text, page_number):
                part["partIndex"] = len(parts)
                parts.append(part)
        return parts
    except Exception as error:
        raise DocumentParseError("SESSION_DOCUMENT_PARSE_FAILED", "PDF 文档解析失败", 422) from error


# 从 DOCX 段落和表格中提取可检索文本。
def _parse_docx(buffer: bytes) -> list[dict[str, Any]]:
    try:
        document = DocxDocument(io.BytesIO(buffer))
        sections = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    sections.append(" | ".join(cells))
        return _build_parts("\n\n".join(sections))
    except Exception as error:
        raise DocumentParseError("SESSION_DOCUMENT_PARSE_FAILED", "DOCX 文档解析失败", 422) from error


# 解析上传文档，返回与两套 Agent API 共用的会话文档结构。
def parse_document(file_content: bytes, filename: str, mime_type: str = "application/octet-stream") -> dict[str, Any]:
    if len(file_content) > SESSION_DOCUMENT_MAX_BYTES:
        raise DocumentParseError("SESSION_DOCUMENT_TOO_LARGE", "临时文档不能超过 2.5MB", 413)

    suffix = Path(filename).suffix.lower()
    if suffix not in SESSION_DOCUMENT_EXTENSIONS:
        raise DocumentParseError(
            "SESSION_DOCUMENT_UNSUPPORTED",
            "临时模式仅支持 txt、md、markdown、json、csv、pdf、docx 文档",
            415,
        )

    if suffix == ".pdf":
        parts = _parse_pdf(file_content)
    elif suffix == ".docx":
        parts = _parse_docx(file_content)
    else:
        try:
            text = file_content.decode("utf-8").lstrip("\ufeff")
        except UnicodeDecodeError as error:
            raise DocumentParseError("SESSION_DOCUMENT_INVALID_ENCODING", "文档必须使用 UTF-8 编码", 422) from error
        parts = _build_parts(text, 1)

    extracted_size = len("\n\n".join(str(part["content"]) for part in parts).encode("utf-8"))
    if extracted_size > SESSION_DOCUMENT_MAX_TEXT_BYTES:
        raise DocumentParseError("SESSION_DOCUMENT_TEXT_TOO_LARGE", "文档正文不能超过 1MB", 413)
    if not parts:
        raise DocumentParseError("SESSION_DOCUMENT_EMPTY", "文档内容不能为空", 400)

    return {
        "localId": f"upload-{uuid.uuid4()}",
        "filename": filename[:255],
        "mimeType": mime_type or "application/octet-stream",
        "size": len(file_content),
        "parserVersion": "session-document-v2",
        "parts": parts,
    }


# 将解析后的文档段落拼接为知识库上传正文。
def document_text(document: dict[str, Any]) -> str:
    return "\n\n".join(
        str(part.get("content") or "").strip()
        for part in document.get("parts", [])
        if str(part.get("content") or "").strip()
    )
